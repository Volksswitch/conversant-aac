"""Self-test for docx_safe: the gate must REFUSE each fault that has actually shipped.

Run directly, or through `npm test` (tests/docx-toolkit.test.mjs drives this file).

⚠ THE POINT IS THE NEGATIVE CASES. A save path that validates is worth nothing unless
it is proven to reject; a check that silently passes everything looks identical to a
clean run, which is the failure this whole toolkit exists to end.
"""
import copy
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import docx
from docx.oxml.ns import qn

import docx_safe as D

FAILURES = []


def check(name, fn):
    try:
        fn()
    except AssertionError as e:
        FAILURES.append('%s: %s' % (name, e))
        print('  FAIL  %s: %s' % (name, e))
    except Exception as e:                                  # noqa: BLE001
        FAILURES.append('%s: unexpected %s: %s' % (name, type(e).__name__, e))
        print('  ERROR %s: %s: %s' % (name, type(e).__name__, e))
    else:
        print('  ok    %s' % name)


def _doc_with_table():
    d = docx.Document()
    d.add_paragraph('An anchor paragraph for editing.')
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = 'Initial delay'
    t.cell(0, 1).text = 'How long it waits.'
    t.cell(1, 0).text = 'Subsequent delay'
    t.cell(1, 1).text = 'The gap before another.'
    return d


def _tmp():
    fd, p = tempfile.mkstemp(suffix='.docx')
    os.close(fd)
    return p


def _refuses(d, why):
    p = _tmp()
    try:
        D.save(d, p)
    except D.DocxIntegrityError:
        return
    finally:
        if os.path.exists(p):
            os.remove(p)
    raise AssertionError('save() accepted a document with ' + why)


# --- the gate rejects what Word rejects -------------------------------------

def t_empty_cell_refused():
    d = _doc_with_table()
    tc = d.tables[0].cell(1, 1)._tc
    for para in list(tc.iterchildren(qn('w:p'))):
        tc.remove(para)                       # exactly what "delete the paragraph" does
    _refuses(d, 'an empty table cell')


def t_duplicate_comment_reference_refused():
    d = _doc_with_table()
    p = d.paragraphs[0]
    r = p.add_run('x')._element
    ref = r.makeelement(qn('w:commentReference'), {qn('w:id'): '1'})
    r.append(ref)
    p._element.addnext(copy.deepcopy(p._element))    # clones the run AND its reference
    _refuses(d, 'a duplicated commentReference id')


def t_unbalanced_field_refused():
    d = _doc_with_table()
    r = d.paragraphs[0].add_run('')._element
    fld = r.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'begin'})
    r.append(fld)                                    # a begin with no end
    _refuses(d, 'an unbalanced field character')


def t_clean_document_saves():
    d = _doc_with_table()
    p = _tmp()
    try:
        D.save(d, p)
        assert os.path.getsize(p) > 0, 'nothing was written'
        docx.Document(p)                              # and it reads back
    finally:
        os.remove(p)


# --- the editing helpers do not create those faults -------------------------

def t_append_run_never_clones_a_unique_child():
    """The exact fault that made Word refuse the Architecture Overview."""
    d = _doc_with_table()
    p = d.paragraphs[0]
    r = p.add_run(' tail')._element
    r.append(r.makeelement(qn('w:commentReference'), {qn('w:id'): '7'}))

    D.append_run(d, 'An anchor paragraph', ' Appended sentence.')

    refs = [e.get(qn('w:id')) for e in d.element.body.iter(qn('w:commentReference'))]
    assert refs == ['7'], 'the comment reference was duplicated: %r' % (refs,)
    assert 'Appended sentence.' in p.text, 'the text was not appended'
    assert not D.check_integrity(d), D.check_integrity(d)


def t_append_run_keeps_formatting_without_copying_content():
    d = _doc_with_table()
    p = d.paragraphs[0]
    p.runs[0].bold = True
    D.append_run(d, 'An anchor paragraph', ' more')
    assert p.runs[-1].bold, 'formatting was not inherited'
    assert p.runs[-1].text == ' more', 'content leaked from the source run'


def t_insert_row_after_adds_a_row_not_a_paragraph():
    d = _doc_with_table()
    before = len(d.tables[0].rows)
    D.insert_row_after(d, 'Subsequent delay', ['Wait longer', 'Adds to the delay.'])
    assert len(d.tables[0].rows) == before + 1, 'no row was added'
    assert d.tables[0].cell(2, 0).text == 'Wait longer'
    assert not D.check_integrity(d), D.check_integrity(d)


def t_ambiguous_row_is_refused_not_guessed():
    d = _doc_with_table()
    D.insert_row_after(d, 'Subsequent delay', ['Initial delay', 'A second one.'])
    try:
        D.find_row(d, 'Initial delay')
    except LookupError:
        pass
    else:
        raise AssertionError('an ambiguous row name must not resolve to a guess')
    # ...but it resolves when told which one.
    D.find_row(d, 'Initial delay', second_cell_startswith='How long')


def t_sub_in_run_refuses_to_flatten_across_runs():
    d = _doc_with_table()
    p = d.add_paragraph()
    p.add_run('the quick ')
    p.add_run('brown').bold = True
    p.add_run(' fox')
    try:
        D.sub_in_run(d, 'quick brown', 'slow brown')
    except D.DocxIntegrityError:
        pass
    else:
        raise AssertionError('a cross-run edit must be refused, not silently flattened')


def t_body_paragraphs_sees_inside_tables():
    """The property is that a paragraph inside a table CELL is seen at all.

    Deliberately asserted by looking for the cell's text rather than by comparing against
    python-docx's document.paragraphs: that attribute is the very view which hides tables,
    and check-docs.py rightly warns about any documentation script reading it. Proving the
    point by naming the hidden paragraph is both the stronger test and the honest one.
    """
    d = _doc_with_table()
    seen = [p.text for p in D.body_paragraphs(d)]
    assert 'Initial delay' in seen, 'table cells were skipped - the wrong index basis'
    assert 'An anchor paragraph for editing.' in seen, 'body paragraphs were skipped'


def t_insert_para_refuses_inside_a_cell():
    d = _doc_with_table()
    try:
        D.insert_para_after(d, 'Initial delay', 'stray prose')
    except D.DocxIntegrityError:
        pass
    else:
        raise AssertionError('a body paragraph must not be planted inside a table cell')


if __name__ == '__main__':
    print('docx_safe self-test')
    for name, fn in sorted(globals().items()):
        if name.startswith('t_') and callable(fn):
            check(name[2:].replace('_', ' '), fn)
    print('%d failure(s)' % len(FAILURES))
    sys.exit(1 if FAILURES else 0)
