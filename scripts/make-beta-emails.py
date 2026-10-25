"""Make the Word copies of the beta emails from BETA-EMAILS.md.

    python scripts/make-beta-emails.py

BETA-EMAILS.md is the master; this rewrites every .docx in Other/Beta Testing/Emails.
Each "## <id>. <name>" section becomes one file. The lines before the quoted email
are the when/to notes; the quoted lines are the email. Markdown handled: **bold**,
*italic*, [text](url) as a clickable link, "- " bullets and "1. " numbered items.
"""
import os, re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
SRC = os.path.join(ROOT, 'BETA-EMAILS.md')
OUT = os.path.join(ROOT, 'Other', 'Beta Testing', 'Emails')

TOKEN = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))')


def add_link(p, text, url, bold=False):
    rid = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rpr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
    if bold:
        rpr.append(OxmlElement('w:b'))
    r.append(rpr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
    h.append(r); p._p.append(h)


def add_runs(p, s):
    for part in TOKEN.split(s):
        if not part:
            continue
        m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
        if m and not m.group(1).startswith(('first', 'helper', 'their', 'referrer', 'Referrer', 'AI key', 'a Windows', 'an iPad')):
            add_link(p, m.group(1), m.group(2))
        elif part.startswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('*'):
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)


def build(num, title, lines):
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    doc.add_heading(f'Beta email {num}: {title}', level=1)

    notes, body = [], []
    for ln in lines:
        if ln.startswith('>'):
            body.append(ln[2:] if ln.startswith('> ') else ln[1:])
        elif ln.strip() and not body:
            if ln.startswith('**') or not notes:
                notes.append(ln.strip())
            else:
                notes[-1] += ' ' + ln.strip()
    for t in notes:
        add_runs(doc.add_paragraph(), t)
    doc.add_paragraph()

    buf, kind = [], None

    def flush():
        nonlocal buf, kind
        if buf:
            style = {'li': 'List Bullet', 'num': 'List Number'}.get(kind)
            text = ' '.join(buf)
            if kind == 'num':
                text = re.sub(r'^\d+\.\s+', '', text)
            p = doc.add_paragraph(style=style)
            if kind == 'indent':
                p.paragraph_format.left_indent = Inches(0.5)
            add_runs(p, text)
        buf, kind = [], None

    for raw in body:
        ln = raw.strip()
        indented = raw.startswith('   ')
        if not ln:
            flush(); continue
        if ln.startswith('- '):
            flush(); kind = 'li'; buf = [ln[2:]]
        elif re.match(r'\d+\. ', ln):
            flush(); kind = 'num'; buf = [ln]
        elif ln.startswith('Subject:'):
            flush(); add_runs(doc.add_paragraph(), '**' + ln + '**')
        elif buf and buf[-1].endswith(',') and len(buf[-1]) < 40:
            flush(); buf = [ln]
        elif buf:
            buf.append(ln)
        else:
            # a paragraph indented under a numbered step stays indented under it
            buf = [ln]; kind = 'indent' if indented else None
    flush()

    path = os.path.join(OUT, f'Beta email {num} - {title}.docx')
    doc.save(path)
    return path


def main():
    os.makedirs(OUT, exist_ok=True)
    for old in os.listdir(OUT):
        if old.startswith('Beta email ') and old.endswith('.docx'):
            try:
                os.remove(os.path.join(OUT, old))
            except PermissionError:
                print(f'Could not remove "{old}" - it is open, probably in Word. Close it and run this again.')
    text = open(SRC, encoding='utf-8').read()
    for sec in re.split(r'^## ', text, flags=re.M)[1:]:
        head, *lines = sec.splitlines()
        m = re.match(r'(\d+[a-z]?)\. (.+)', head)
        if not m or not any(l.startswith('>') for l in lines):
            continue
        title = m.group(2).replace(' - ', ', ').replace('The kind decline', 'Decline')
        print(build(m.group(1), title, lines))


main()
